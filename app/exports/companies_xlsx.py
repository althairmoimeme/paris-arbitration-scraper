"""Export the Clients potentiels (companies with signals) to XLSX."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from app.config import EXPORT_DIR
from app.database import (
    Company,
    CompanySignal,
    ICSIDCase,
    SessionLocal,
)
from app.logging_setup import setup as _setup_logging

_setup_logging()
from loguru import logger  # noqa: E402


HEADER_FILL = PatternFill(start_color="0F1B2D", end_color="0F1B2D", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
BORDER = Border(
    left=Side(style="thin", color="DDDDDD"),
    right=Side(style="thin", color="DDDDDD"),
    top=Side(style="thin", color="DDDDDD"),
    bottom=Side(style="thin", color="DDDDDD"),
)
FILL_RED = PatternFill(start_color="E63323", end_color="E63323", fill_type="solid")
FILL_YELLOW = PatternFill(start_color="FFF4C2", end_color="FFF4C2", fill_type="solid")
FILL_GREEN = PatternFill(start_color="C8F7C5", end_color="C8F7C5", fill_type="solid")


COLUMNS = [
    ("Société", 32),
    ("Secteur", 14),
    ("Sous-secteur", 28),
    ("Pays respondent (si ICSID)", 18),
    ("Exposure Afrique", 10),
    ("Nombre de signaux", 10),
    ("Hot score", 10),
    ("Site web", 28),
    ("ICSID cases", 60),
    ("Marchés publics (DECP) cumulé €", 18),
    ("Signaux (détail)", 80),
    ("Angle commercial", 80),
]


def _format_amount(n):
    if n is None or n == 0:
        return "n/c"
    if n >= 1_000_000_000:
        return f"{n / 1_000_000_000:.1f} Md€"
    if n >= 1_000_000:
        return f"{n / 1_000_000:.1f} M€"
    return f"{n:,.0f} €".replace(",", " ")


def export_companies_xlsx(path: Path | None = None) -> Path:
    path = path or (EXPORT_DIR / "clients_potentiels.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Clients potentiels"

    # Header
    for i, (name, w) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

    with SessionLocal() as session:
        companies = session.scalars(
            select(Company).order_by(
                Company.hot_score.desc(), Company.signal_count.desc(),
            )
        ).all()
        r = 2
        for company in companies:
            icsid_cases = session.scalars(
                select(ICSIDCase).where(ICSIDCase.company_id == company.id)
            ).all()
            signals = session.scalars(
                select(CompanySignal).where(CompanySignal.company_id == company.id)
            ).all()

            icsid_summary = " ; ".join(
                f"{c.case_number} vs {c.respondent_state or 'n/c'}"
                f" ({c.economic_sector or 'n/c'}, {c.status or 'n/c'})"
                for c in icsid_cases[:5]
            )
            decp_total = sum((s.amount_eur or 0) for s in signals if s.kind == "decp_contract")
            signals_text = " || ".join(
                f"[{s.kind}] {s.title} · {(s.detail or '')[:120]}"
                for s in signals[:10]
            )
            respondent_countries = ", ".join(
                sorted({c.respondent_state for c in icsid_cases if c.respondent_state})
            )

            vals = [
                company.name,
                (company.sector or "").replace("_", " ").title(),
                company.sector_detail or "",
                respondent_countries,
                "YES" if company.has_africa_exposure else "",
                company.signal_count or 0,
                company.hot_score or 0,
                company.website or "",
                icsid_summary,
                _format_amount(decp_total),
                signals_text,
                company.commercial_angle or "",
            ]
            for i, v in enumerate(vals, start=1):
                c = ws.cell(row=r, column=i, value=v)
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.border = BORDER
            # Highlight Africa exposure column
            if company.has_africa_exposure:
                ws.cell(row=r, column=5).fill = FILL_GREEN
                ws.cell(row=r, column=5).font = Font(bold=True)
            # Highlight hot score
            if (company.hot_score or 0) >= 80:
                ws.cell(row=r, column=7).fill = FILL_RED
                ws.cell(row=r, column=7).font = Font(bold=True, color="FFFFFF")
            elif (company.hot_score or 0) >= 50:
                ws.cell(row=r, column=7).fill = FILL_YELLOW
            ws.row_dimensions[r].height = 60
            r += 1

    last_col = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A1:{last_col}{r - 1}"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    logger.info(f"Wrote {path.name}  ({r - 2} companies)")
    return path
